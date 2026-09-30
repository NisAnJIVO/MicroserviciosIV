# -*- coding: utf-8 -*-
import docx

def fill_doc():
    doc = docx.Document("Practica 6. Mensajeria Asincrona con RabbitMQ - Hoja de Evidencias.docx")

    # 1. Tabla 0: Datos del estudiante y entorno
    t0 = doc.tables[0]
    t0.rows[4].cells[1].text = "Windows 11 Pro 64-bit | Node.js v24.14.0 | Docker Engine v29.2.1 | Imagen rabbitmq:3.13-management"
    t0.rows[5].cells[1].text = "https://github.com/NisAnJIVO/MicroserviciosIV.git"

    # 2. Tabla 2: Control de avance (Marcar con X los laboratorios)
    t2 = doc.tables[2]
    for row in t2.rows[1:]:
        row.cells[3].text = "X"

    # 3. Pregunta 1 -> Tabla 6
    doc.tables[6].rows[0].cells[0].text = (
        "• ¿Qué gana el sistema al introducir el broker (RabbitMQ)?\n"
        "  1. Desacoplamiento temporal y espacial: El emisor (productor) no necesita conocer la ubicación física, "
        "dirección IP ni el estado de disponibilidad del receptor (consumidor), ni requiere que ambos estén activos "
        "simultáneamente para procesar una operación.\n"
        "  2. Amortiguación y nivelación de carga (Load Leveling / Buffering): Ante picos masivos de peticiones, los "
        "mensajes se acumulan en colas durables, permitiendo a los servicios consumidores procesar a un ritmo constante "
        "sin colapsar la infraestructura.\n"
        "  3. Resiliencia y latencia mínima percibida: El productor concluye transacciones de cara al usuario en milisegundos "
        "sin quedar bloqueado ni heredar la latencia o fallos de servicios dependientes secundarios.\n\n"
        "• ¿Qué problema nuevo se crea con él?\n"
        "  1. Complejidad operativa y punto único de fallo (SPOF): Se introduce una pieza central de infraestructura "
        "que debe administrarse, monitorearse, respaldarse y configurarse en clúster para garantizar alta disponibilidad.\n"
        "  2. Pérdida de inmediatez y consistencia eventual: Se abandona la consistencia inmediata (ACID) por consistencia eventual. "
        "Si un procesamiento asíncrono posterior falla definitivamente, el sistema requiere mecanismos de compensación complejos "
        "(patrón Saga, DLQ o intervención manual) en lugar de una respuesta síncrona simple de error."
    )

    # 4. Pregunta 2 -> Tabla 10
    doc.tables[10].rows[0].cells[0].text = (
        "• ¿Qué hay que modificar en el productor si el área de finanzas también necesita enterarse?\n"
        "  Absolutamente NADA en el código ni en la configuración del microservicio productor.\n\n"
        "• Justificación basada en la experiencia práctica:\n"
        "  En el modelo AMQP, el productor está completamente desacoplado de las colas receptoras; su única responsabilidad "
        "es publicar el evento en el exchange ('inscripciones') con una clave de enrutamiento ('inscripcion.confirmada'). "
        "El enrutamiento es competencia exclusiva del broker. Para integrar al área de finanzas, solo se necesita crear una "
        "nueva cola (por ejemplo, 'finanzas.cobros') y declararle un binding hacia el exchange 'inscripciones' con la misma clave "
        "'inscripcion.confirmada'. El exchange copiará de manera transparente y automática cada mensaje entrante tanto a "
        "'notificaciones.correo' como a la cola de finanzas, cumpliendo el principio Open/Closed sin tocar una sola línea del productor."
    )

    # 5. Pregunta 3 -> Tabla 13
    doc.tables[13].rows[0].cells[0].text = (
        "• ¿De qué trabajo se liberó el servicio de Inscripciones?\n"
        "  Se liberó de las operaciones accesorias de entrada/salida (I/O) que no forman parte del núcleo transaccional inmediato: "
        "establecer conexiones síncronas con servidores SMTP o APIs externas de correo, lidiar con reintentos de red, gestionar "
        "timeouts y retener la conexión HTTP del estudiante en espera durante segundos innecesarios.\n\n"
        "• ¿Qué promesa está haciendo a cambio al estudiante que se inscribió?\n"
        "  Le garantiza una doble promesa:\n"
        "  1. Certeza transaccional inmediata: Su inscripción quedó confirmada, registrada con identificador único y cupo reservado de manera persistente en el sistema.\n"
        "  2. Garantía de entrega eventual: El correo de bienvenida o comprobante formal le llegará en un instante posterior tan pronto "
        "el servicio de notificaciones procese la cola, garantizando que el mensaje no se perderá aunque el consumidor se encuentre "
        "temporalmente inactivo o congestionado."
    )

    # 6. Pregunta 4 -> Tabla 16
    doc.tables[16].rows[0].cells[0].text = (
        "• ¿Dónde está en este momento el trabajo pendiente?\n"
        "  El trabajo pendiente se encuentra resguardado en el broker RabbitMQ, almacenado en el búfer de la cola durable "
        "'notificaciones.correo' en estado 'Ready' (listo para ser consumido).\n\n"
        "• ¿Qué recursos está consumiendo mientras espera?\n"
        "  1. Memoria RAM: RabbitMQ mantiene los mensajes en memoria principal para entrega de ultra baja latencia.\n"
        "  2. Almacenamiento en disco (Paging / Disk I/O): Al haber sido publicados con 'persistent: true' (o al superar el umbral "
        "de paginación de memoria), RabbitMQ los escribe en disco en su base de datos interna (/var/lib/rabbitmq).\n"
        "  3. Estructuras de metadatos: Punteros de índice y descriptores de archivo dentro del runtime Erlang del broker.\n\n"
        "• ¿Qué ocurriría si el consumidor no volviera nunca?\n"
        "  La cola continuaría acumulando mensajes indefinidamente con cada nueva inscripción. Al aproximarse a los límites configurados "
        "(vm_memory_high_watermark o disco lleno), RabbitMQ entraría en estado de alarma de flujo (flow control), bloqueando las "
        "conexiones entrantes de todos los productores para evitar la caída del nodo. Eventualmente, se saturaría el disco duro "
        "ocasionando una denegación de servicio (DoS) global si no se tienen políticas de caducidad (TTL) o límites máximos de cola (max-length)."
    )

    # 7. Pregunta 5 -> Tabla 20
    doc.tables[20].rows[0].cells[0].text = (
        "• ¿Por qué 'certificado.emitido' llegó solo a una de las tres colas?\n"
        "  Porque en el exchange temático ('eventos.academicos'), las reglas de binding definidas fueron:\n"
        "  - 'correo.bienvenida' escucha únicamente 'inscripcion.*' (requiere que empiece con la palabra 'inscripcion').\n"
        "  - 'finanzas.cobros' escucha únicamente 'pago.*' (requiere que empiece con la palabra 'pago').\n"
        "  - 'auditoria.todo' escucha '#' (captura todo sin importar palabras ni niveles).\n"
        "  La clave 'certificado.emitido' no coincide con 'inscripcion' ni con 'pago', pero coincide con '#', por lo que fue descartada "
        "por las dos primeras y copiada exclusivamente a la cola de auditoría.\n\n"
        "• Diferencia entre el comodín '*' y '#':\n"
        "  - El asterisco (*) sustituye EXACTAMENTE UNA palabra delimitada por puntos.\n"
        "  - El numeral (#) sustituye CERO, UNA O MÚLTIPLES palabras delimitadas por puntos.\n\n"
        "• Ejemplo en nuestro Proyecto Integrador (Plataforma de Comercio Electrónico y Gestión de Pedidos):\n"
        "  Supongamos eventos de pedidos como 'pedidos.nacional.creado' y 'pedidos.internacional.express.creado':\n"
        "  - El patrón 'pedidos.*.creado' coincidirá con 'pedidos.nacional.creado' (1 palabra intermedia), pero fallará con "
        "'pedidos.internacional.express.creado' (tiene 2 palabras intermedias).\n"
        "  - El patrón 'pedidos.#' coincidirá con absolutamente todos los eventos generados por el dominio de pedidos "
        "(como 'pedidos.creado', 'pedidos.nacional.cancelado' o 'pedidos.internacional.express.pagado'), ideal para un microservicio de analítica global."
    )

    # 8. Pregunta 6 -> Tabla 24
    doc.tables[24].rows[0].cells[0].text = (
        "• ¿Qué pasa exactamente si se cumple solo una de las dos condiciones?\n"
        "  1. Cola Durable pero Mensaje Transitorio (persistent: false): Al reiniciar el broker, RabbitMQ recrea la definición y estructura "
        "de la cola, pero todos los mensajes que contenía desaparecen por completo porque solo residían en la memoria RAM.\n"
        "  2. Mensaje Persistente (persistent: true) pero Cola Transitoria (durable: false): Al reiniciar el broker, la cola completa se destruye. "
        "Al eliminarse la cola que contenía los mensajes, RabbitMQ purga automáticamente todos sus mensajes asociados, perdiéndose los datos a pesar de haber sido escritos en disco.\n"
        "  Conclusión: Se requiere OBLIGATORIAMENTE que ambos factores sean verdaderos para garantizar supervivencia tras un reinicio.\n\n"
        "• ¿Por qué la durabilidad tiene un costo que no siempre conviene pagar?\n"
        "  Porque la persistencia exige sincronizaciones forzadas y continuas a disco magnético o SSD (llamadas fsync del sistema operativo) "
        "y mantenimiento de bitácoras transaccionales de escritura anticipada (WAL). Esto introduce latencia de I/O y reduce el throughput "
        "de decenas de miles de mensajes por segundo en memoria a solo unos pocos miles en disco.\n"
        "  No conviene pagar este costo en casos donde el volumen es masivo y la pérdida aislada de eventos no afecta al negocio, como en "
        "telemetría de sensores IoT, coordenadas GPS de vehículos en tiempo real, métricas de monitoreo de servidores o registros de clics web (clickstream)."
    )

    # 9. Pregunta 7 -> Tabla 27
    doc.tables[27].rows[0].cells[0].text = (
        "• ¿Por qué sería un error ponerle un consumidor automático que reintente sin más?\n"
        "  Porque si un mensaje aterrizó en la cola de mensajes muertos (DLQ) tras agotar todos sus reintentos, significa que su falla es "
        "determinista o sistemática (un 'poison message' o mensaje envenenado: un JSON malformado, un correo inválido como 'ana.punto.usfx.bo', "
        "un ID que viola integridad referencial en la base de datos o un tipo de dato erróneo). Un consumidor que reintente automáticamente a ciegas "
        "lo reinyectará en la cola principal, provocando un bucle infinito de caídas, saturación inútil de CPU y memoria, y contaminación de bitácoras sin resolver el problema.\n\n"
        "• Procedimiento paso a paso para atender la cola muerta:\n"
        "  1. Detección y Notificación: El sistema de observabilidad (Prometheus/Grafana) detecta mensajes en 'notificaciones.muertos' y dispara una alerta de soporte.\n"
        "  2. Diagnóstico e Inspección: El operador inspecciona el mensaje con 'Get messages' o CLI, examinando el payload y las cabeceras "
        "(x-death, excepciones registradas y marcas de tiempo) para aislar la causa raíz del fallo.\n"
        "  3. Corrección de la Causa Raíz:\n"
        "     - Si se debe a un bug en el microservicio o dependencia caída, se despliega la corrección o se reactiva el servicio externo.\n"
        "     - Si se debe a datos corruptos del usuario (ej. correo mal escrito), se edita y sanea manualmente el mensaje en el registro transaccional.\n"
        "  4. Reenrutamiento Seguro (Reprocesamiento): Mediante una herramienta como rabbitmqadmin, un script Shovel o consola, se reenvía el "
        "mensaje corregido al exchange correspondiente para que se procese normalmente.\n"
        "  5. Mejora Preventiva: Se registra el caso y se agrega una regla de validación preventiva en el productor (ej. validación formal de correo) "
        "para que errores de ese tipo se rechacen con HTTP 400 antes de ingresar a la cola."
    )

    # 10. Pregunta 8 -> Tabla 30
    doc.tables[30].rows[0].cells[0].text = (
        "• ¿Por qué le resulta imposible al broker garantizar la entrega 'exactamente una vez'?\n"
        "  Debido a las limitaciones inherentes a las redes en sistemas distribuidos (el Teorema CAP y el Problema de los Dos Generales). "
        "Cuando el consumidor termina de procesar un mensaje y envía la confirmación (ack) por la red, la conexión puede sufrir un corte, "
        "un timeout o el consumidor puede caerse justo antes de que el paquete de red alcance al broker. Al cerrarse la conexión sin confirmación, "
        "RabbitMQ no puede distinguir si el procesamiento se completó o si falló; por ende, para cumplir su promesa de 'al menos una vez' (no perder mensajes), "
        "está obligado a reencolarlo y entregarlo a otro consumidor, provocando una entrega duplicada.\n\n"
        "• ¿En qué parte del sistema queda la responsabilidad?\n"
        "  La responsabilidad recae enteramente en el MICROSERVICIO CONSUMIDOR. Debe implementarse lógica de IDEMPOTENCIA: registrar en un "
        "almacén persistente (tabla con clave única o índice) el identificador del evento ('messageId'). Si llega un evento repetido, se reconoce e ignora "
        "el efecto secundario, confirmando el mensaje con ack para que no siga circulando.\n\n"
        "• Ejemplo y daño visible en nuestro Proyecto Integrador (Plataforma de Comercio Electrónico y Pedidos):\n"
        "  En el microservicio de Pagos y Facturación, al recibir el evento 'PedidoConfirmado': si el consumidor debita el saldo de la tarjeta o cuenta "
        "del cliente y sufre una pérdida de red antes de enviar el ack, el broker reentregará el evento. Si el consumidor no fuera idempotente, "
        "cobraría el dinero POR SEGUNDA VEZ, generando un cobro duplicado y daño financiero directo y visible en el extracto bancario del usuario."
    )

    # 11. Cierre de la Parte 1 -> Tablas 31, 32, 33
    doc.tables[31].rows[0].cells[0].text = (
        "La configuración y comprensión del circuito de Dead Letter Queue (DLQ) mediante el encabezado x-death y la combinación con colas de "
        "espera basadas en TTL. Articular exchanges directos de reintento, colas con tiempo de expiración y el salto ordenado hacia la cola muerta "
        "final exigió comprender con exactitud el ciclo de vida de los mensajes en el protocolo AMQP."
    )

    doc.tables[32].rows[0].cells[0].text = (
        "El error PRECONDITION_FAILED - inequivalent arg al intentar redeclarar la cola 'notificaciones.correo' agregando el argumento x-dead-letter-exchange. "
        "Se resolvió comprendiendo que RabbitMQ prohíbe por diseño mutar la estructura de una cola existente en caliente para proteger la consistencia de los datos, "
        "por lo que fue necesario eliminarla limpiamente desde la consola de administración antes de ejecutar el script dlq.js."
    )

    doc.tables[33].rows[0].cells[0].text = (
        "Adoptaría desde el inicio una convención formal de nombres de tópicos jerárquicos y persistencia transaccional (Transactional Outbox Pattern) "
        "respaldada por una base de datos relacional para el registro de idempotencia en lugar de estructuras en memoria, garantizando consistencia absoluta "
        "entre los estados del microservicio y los eventos emitidos al broker."
    )

    doc.save("Practica 6. Mensajeria Asincrona con RabbitMQ - Hoja de Evidencias.docx")
    print("Documento Word completado y guardado exitosamente!")

if __name__ == "__main__":
    fill_doc()
